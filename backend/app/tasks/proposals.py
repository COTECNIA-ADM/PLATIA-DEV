"""Paquete de documentos de una propuesta aprobada (cola ``documents``, Capítulo 9).

En esta fase se genera el DOCX editable. El PDF y el PPTX llegan con la
generación de documentos de la fase 5.
"""
import io

import structlog
from docx import Document as DocxDocument
from docx.shared import Pt

from ..extensions import db
from ..models import Proposal, ProposalAttachment
from ..queue import PlatformTask, celery
from ..services import storage
from ..services.notifications import notify

log = structlog.get_logger(__name__)


def render_docx(proposal: Proposal) -> bytes:
    doc = DocxDocument()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(proposal.title, level=0)
    opp = proposal.opportunity
    meta = [("Oportunidad", opp.title if opp else ""),
            ("Cliente", opp.client.name if opp and opp.client else "—"),
            ("Tipo de propuesta", proposal.proposal_type),
            ("Versión", str(proposal.version_number))]
    if proposal.total_amount_eur is not None:
        meta.append(("Importe total", f"{proposal.total_amount_eur:,.2f} €".replace(",", "X")
                     .replace(".", ",").replace("X", ".")))
    table = doc.add_table(rows=0, cols=2)
    for label, value in meta:
        row = table.add_row().cells
        row[0].text, row[1].text = label, value
    for section in sorted((s for s in proposal.sections if s.is_active), key=lambda s: s.position):
        doc.add_heading(section.title, level=1)
        for paragraph in (section.final_content or "").split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@celery.task(name="tasks.proposals.generate_package", base=PlatformTask, bind=True, max_retries=3)
def generate_package(self, proposal_id: int) -> dict:
    proposal = db.session.get(Proposal, proposal_id)
    if proposal is None or not proposal.is_active:
        return {"proposal_id": proposal_id, "status": "missing"}
    try:
        content = render_docx(proposal)
    except Exception as exc:  # noqa: BLE001
        raise self.retry_later(exc)
    from werkzeug.datastructures import FileStorage

    filename = f"propuesta-{proposal.id}-v{proposal.version_number}.docx"
    stored = storage.save_upload(FileStorage(io.BytesIO(content), filename=filename), subdir="proposals")
    # Solo se conserva el paquete generado más reciente
    for old in proposal.attachments:
        if old.kind == "generated" and old.is_active:
            old.soft_delete()
    attachment = ProposalAttachment(kind="generated", filename=filename, mime_type=stored.mime_type,
                                    storage_path=stored.storage_path, file_size_bytes=stored.size,
                                    file_hash=stored.sha256)
    proposal.attachments.append(attachment)
    db.session.commit()
    recipient = proposal.owner_user_id or proposal.created_by
    if recipient:
        notify(recipient, "REPORT_READY", "Paquete de propuesta disponible",
               f"El documento de «{proposal.title[:200]}» está listo para descargar.",
               link=f"/proposals/{proposal.id}", extra_data={"formats": ["docx"]})
    log.info("proposal_package_generated", proposal_id=proposal.id, attachment_id=attachment.id)
    return {"proposal_id": proposal.id, "attachment_id": attachment.id, "formats": ["docx"]}
