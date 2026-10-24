"""Importación de competidores y objetivos desde Excel (Capítulo 23).

El fichero tiene dos hojas (las dos opcionales): ``competidores`` y ``objetivos``.
La primera fila de cada hoja son los encabezados; el orden de las columnas da
igual. Se procesa en dos pasos:

1. **Comprobación** (``dry_run``): se valida cada fila y se dice qué se crearía y
   qué se actualizaría, sin tocar nada.
2. **Importación**: solo si no hay ningún error en ninguna hoja. Todo o nada, en
   una transacción, y queda en la auditoría.

Los competidores se identifican por el nombre (sin distinguir mayúsculas) y los
objetivos por métrica, mes y línea. Se procesa en la petición, no en Celery: el
fichero está limitado a ``MAX_ROWS`` filas por hoja y se lee en modo solo lectura.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import structlog
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..errors import ValidationError
from ..extensions import db
from ..models import BusinessTarget, CompetitorProfile, ServiceLine

log = structlog.get_logger(__name__)

MAX_ROWS = 2000
MAX_CELL = 2000
LIST_SPLIT = re.compile(r"\s*(?:;|\n)\s*")
PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

TYPE_LABELS = {"directo": "direct", "direct": "direct", "indirecto": "indirect", "indirect": "indirect",
               "sustitutivo": "substitute", "substitute": "substitute"}
PRICE_LABELS = {"-1": -1, "0": 0, "1": 1, "más barato": -1, "mas barato": -1, "similar": 0, "más caro": 1,
                "mas caro": 1}
METRIC_LABELS = {"nuevas": "new_count", "new_count": "new_count", "importe nuevas": "new_eur",
                 "new_eur": "new_eur", "ganadas": "won_count", "won_count": "won_count",
                 "importe ganado": "won_eur", "won_eur": "won_eur"}

COMPETITOR_COLUMNS = ["nombre", "tipo", "web", "descripcion", "fortalezas", "debilidades", "oportunidades",
                      "amenazas", "lineas", "sectores", "precio"]
TARGET_COLUMNS = ["periodo", "metrica", "linea", "objetivo"]
SHEETS = {"competidores": COMPETITOR_COLUMNS, "objetivos": TARGET_COLUMNS}


@dataclass
class SheetResult:
    rows: int = 0
    create: int = 0
    update: int = 0
    errors: list[dict] = field(default_factory=list)
    items: list[dict] = field(default_factory=list)  # filas válidas, listas para aplicar

    def to_dict(self) -> dict:
        return {"rows": self.rows, "create": self.create, "update": self.update, "errors": self.errors[:200],
                "error_count": len(self.errors)}


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()[:MAX_CELL]


def _list(value) -> list[str]:
    return [item[:300] for item in LIST_SPLIT.split(_text(value)) if item][:20]


def _read_rows(ws, sheet: str, result: SheetResult) -> list[tuple[int, dict]]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    names = [_text(h).lower() for h in header]
    expected = SHEETS[sheet]
    missing = [c for c in expected[:1] if c not in names]
    if missing:
        result.errors.append({"row": 1, "field": missing[0], "message": f"Falta la columna «{missing[0]}»."})
        return []
    unknown = [n for n in names if n and n not in expected]
    if unknown:
        result.errors.append({"row": 1, "field": unknown[0],
                              "message": f"Columna desconocida «{unknown[0]}». Usa la plantilla."})
        return []
    out = []
    for number, values in enumerate(rows, start=2):
        if all(v is None or _text(v) == "" for v in values):
            continue
        if len(out) >= MAX_ROWS:
            result.errors.append({"row": number, "field": None, "message": f"Más de {MAX_ROWS} filas."})
            break
        out.append((number, {name: values[i] if i < len(values) else None for i, name in enumerate(names) if name}))
    return out


def _check_competitors(rows, lines: set[str], result: SheetResult) -> None:
    existing = {c.name.lower(): c for c in db.session.scalars(db.select(CompetitorProfile))}
    seen = set()
    for number, row in rows:
        result.rows += 1
        name = _text(row.get("nombre"))[:200]

        def error(field: str, message: str, number: int = number):
            result.errors.append({"row": number, "field": field, "message": message})

        if not name:
            error("nombre", "El nombre es obligatorio.")
            continue
        if name.lower() in seen:
            error("nombre", f"«{name}» está repetido en el fichero.")
            continue
        seen.add(name.lower())
        kind = TYPE_LABELS.get(_text(row.get("tipo")).lower() or "directo")
        if kind is None:
            error("tipo", "Tipo no válido: directo, indirecto o sustitutivo.")
            continue
        price = PRICE_LABELS.get(_text(row.get("precio")).lower() or "0")
        if price is None:
            error("precio", "Precio no válido: más barato, similar o más caro (o -1, 0, 1).")
            continue
        service_lines = _list(row.get("lineas"))
        bad = [code for code in service_lines if code not in lines]
        if bad:
            error("lineas", f"Línea de servicio desconocida: {bad[0]}. Usa los códigos de la plataforma.")
            continue
        website = _text(row.get("web"))[:255] or None
        if website and not re.match(r"^https?://", website):
            website = "https://" + website
        current = existing.get(name.lower())
        result.update += current is not None and current.is_active
        result.create += current is None or not current.is_active
        result.items.append({"id": current.id if current else None, "name": name, "competitor_type": kind,
                             "website": website, "description": _text(row.get("descripcion")) or None,
                             "strengths": _list(row.get("fortalezas")), "weaknesses": _list(row.get("debilidades")),
                             "opportunities": _list(row.get("oportunidades")), "threats": _list(row.get("amenazas")),
                             "service_lines": service_lines, "sectors": _list(row.get("sectores")),
                             "price_positioning": price})


def _check_targets(rows, lines: set[str], result: SheetResult) -> None:
    existing = {(t.metric, t.period, t.service_line): t for t in db.session.scalars(db.select(BusinessTarget))}
    seen = set()
    for number, row in rows:
        result.rows += 1

        def error(field: str, message: str, number: int = number):
            result.errors.append({"row": number, "field": field, "message": message})

        raw_period = row.get("periodo")
        period = raw_period.strftime("%Y-%m") if hasattr(raw_period, "strftime") else _text(raw_period)
        if not PERIOD_RE.match(period):
            error("periodo", "El periodo debe ser AAAA-MM (por ejemplo 2026-11).")
            continue
        metric = METRIC_LABELS.get(_text(row.get("metrica")).lower())
        if metric is None:
            error("metrica", "Métrica no válida: nuevas, importe nuevas, ganadas o importe ganado.")
            continue
        line = _text(row.get("linea"))
        if line and line not in lines:
            error("linea", f"Línea de servicio desconocida: {line}.")
            continue
        try:
            target = Decimal(_text(row.get("objetivo")).replace(",", "."))
        except InvalidOperation:
            error("objetivo", "El objetivo debe ser un número.")
            continue
        if target < 0 or target >= Decimal("1e13") or not target.is_finite():
            error("objetivo", "El objetivo debe ser un número positivo.")
            continue
        key = (metric, period, line)
        if key in seen:
            error("periodo", "Objetivo repetido en el fichero (misma métrica, mes y línea).")
            continue
        seen.add(key)
        current = existing.get(key)
        result.update += current is not None and current.is_active
        result.create += current is None or not current.is_active
        result.items.append({"id": current.id if current else None, "metric": metric, "period": period,
                             "service_line": line, "target_value": target.quantize(Decimal("0.01"))})


def process(content: bytes, *, dry_run: bool, user_id: int, can_competitors: bool, can_targets: bool) -> dict:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, KeyError, ValueError, TypeError) as exc:
        raise ValidationError("No se puede leer el fichero: tiene que ser un Excel (.xlsx).") from exc
    lines = set(db.session.scalars(db.select(ServiceLine.code)))
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}
    results: dict[str, SheetResult] = {}
    try:
        if "competidores" in sheets:
            if not can_competitors:
                raise ValidationError("No tienes permiso para importar competidores.")
            results["competidores"] = SheetResult()
            _check_competitors(_read_rows(sheets["competidores"], "competidores", results["competidores"]), lines,
                               results["competidores"])
        if "objetivos" in sheets:
            if not can_targets:
                raise ValidationError("No tienes permiso para importar objetivos.")
            results["objetivos"] = SheetResult()
            _check_targets(_read_rows(sheets["objetivos"], "objetivos", results["objetivos"]), lines,
                           results["objetivos"])
    finally:
        wb.close()
    if not results:
        raise ValidationError("El fichero no tiene ninguna hoja «competidores» ni «objetivos».")

    has_errors = any(r.errors for r in results.values())
    applied = not dry_run and not has_errors
    if applied:
        _apply(results, user_id)
    return {"dry_run": dry_run, "applied": applied, "has_errors": has_errors,
            "sheets": {name: r.to_dict() for name, r in results.items()}}


def _apply(results: dict[str, SheetResult], user_id: int) -> None:
    for item in results.get("competidores", SheetResult()).items:
        obj = db.session.get(CompetitorProfile, item.pop("id")) if item.get("id") else None
        if obj is None:
            db.session.add(CompetitorProfile(**item, created_by=user_id))
            continue
        obj.restore()
        for key, value in item.items():
            if getattr(obj, key) != value:
                setattr(obj, key, value)
    for item in results.get("objetivos", SheetResult()).items:
        obj = db.session.get(BusinessTarget, item.pop("id")) if item.get("id") else None
        if obj is None:
            db.session.add(BusinessTarget(**item, created_by=user_id))
            continue
        obj.restore()
        if obj.target_value != item["target_value"]:
            obj.target_value = item["target_value"]
    db.session.commit()
    log.info("bi_import_applied", **{name: len(r.items) for name, r in results.items()})


def template() -> bytes:
    """Plantilla con las dos hojas, encabezados, una fila de ejemplo y las instrucciones."""
    lines = list(db.session.scalars(db.select(ServiceLine.code).order_by(ServiceLine.code)))
    wb = Workbook()
    ws = wb.active
    ws.title = "competidores"
    ws.append(COMPETITOR_COLUMNS)
    ws.append(["Competidor Ejemplo, S.L.", "directo", "https://ejemplo.com", "Integrador nacional",
               "Precio agresivo; Presencia en sector público", "Poca experiencia en IA", "Alianzas cloud",
               "Entrada de grandes consultoras", ";".join(lines[:2]), "public; health", "más barato"])
    ws2 = wb.create_sheet("objetivos")
    ws2.append(TARGET_COLUMNS)
    ws2.append(["2026-11", "importe ganado", "", 250000])
    ws2.append(["2026-11", "nuevas", lines[0] if lines else "", 12])
    help_ws = wb.create_sheet("instrucciones")
    for row in [
        ["Cómo rellenar esta plantilla"],
        ["Puedes borrar la hoja que no uses. Las filas de ejemplo se pueden sustituir."],
        ["competidores.tipo", "directo, indirecto o sustitutivo"],
        ["competidores.fortalezas/debilidades/oportunidades/amenazas", "varios puntos separados por «;»"],
        ["competidores.lineas", "códigos de línea de servicio separados por «;»: " + ", ".join(lines)],
        ["competidores.precio", "más barato, similar o más caro"],
        ["objetivos.periodo", "mes en formato AAAA-MM"],
        ["objetivos.metrica", "nuevas, importe nuevas, ganadas o importe ganado"],
        ["objetivos.linea", "vacío para el total de la empresa, o un código de línea"],
        ["objetivos.objetivo", "número (importe en euros o número de oportunidades)"],
    ]:
        help_ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
