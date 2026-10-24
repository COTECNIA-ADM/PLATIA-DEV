"""Extracción de texto de documentos (se ejecuta en el worker de la cola documents)."""
from pathlib import Path

MAX_CHARS = 2_000_000  # LONGTEXT admite más, pero no tiene sentido para el análisis


class UnsupportedFormat(Exception):
    pass


def extract_text(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == "docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext in ("txt", "md", "csv"):
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
    else:
        raise UnsupportedFormat(ext)
    return text.replace("\x00", "")[:MAX_CHARS].strip()
