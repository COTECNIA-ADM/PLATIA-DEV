"""Marcado de los documentos generados con IA (fase 8)."""
import io

from docx import Document as DocxDocument

from app.services import rendering

CV = {"name": "Ana Pérez", "headline": "Arquitecta cloud", "summary": "Diez años en AWS.",
      "skills": ["AWS", "Kubernetes"], "experiences": [{"role": "Arquitecta", "company": "ACME"}]}


def test_docx_metadata_and_footer(app):
    doc = DocxDocument(io.BytesIO(rendering.cv_docx(CV)))
    assert doc.core_properties.author == "PLATIA"
    assert "inteligencia artificial" in doc.core_properties.comments
    assert doc.sections[0].footer.paragraphs[0].text == rendering.disclosure()


def test_pdf_metadata_and_footer(app):
    pdf = rendering.cv_pdf(CV)
    assert b"inteligencia artificial" in pdf  # en /Subject
    assert b"/Author (PLATIA)" in pdf


def test_visible_notice_can_be_disabled(app):
    app.config["AI_DISCLOSURE_TEXT"] = ""
    doc = DocxDocument(io.BytesIO(rendering.cv_docx(CV)))
    assert doc.sections[0].footer.paragraphs[0].text == ""
    assert "inteligencia artificial" in doc.core_properties.comments  # los metadatos se mantienen


def test_me_exposes_currency(client, as_role):
    headers = as_role("analyst")
    settings = client.get("/api/auth/me", headers=headers).json["settings"]
    assert settings == {"currency": "EUR", "usd_eur_rate": client.application.config["AI_USD_EUR_RATE"]}
